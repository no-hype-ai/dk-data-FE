"""Tests for EMA EPAR fetcher — parser-error ≠ source-down; loud primary failure.

Feature: PR #415 Phase 1, Task 1.6 (H4, H5)
- H4: corrupt/empty XLSX → parse_error or schema_mismatch (not source_unavailable)
- H5: primary URL failure → ERROR-level log before fallback (not warning-only)

Tests use sync unittest.mock — no asyncio. Mirrors test_ema_regulatory_fetcher.py style.
"""

import io
import logging
from unittest.mock import MagicMock, patch

import openpyxl

from dk_data.ingestion.fetchers.ema_epar import EMAEparFetcher, _EPAR_URLS
from dk_data.ingestion.fetchers.ema_mol import EMAMolFetcher


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _xlsx_bytes(
    header_row_cells: list,
    data_rows: list,
    header_at: int = 8,
) -> bytes:
    """Build a real in-memory XLSX with preamble rows, a header row, then data rows.

    Args:
        header_row_cells: Cell values for the header row (at row index header_at).
        data_rows: List of lists, each is one data row after the header.
        header_at: 0-based row index where the header is placed (default 8).

    Returns:
        Raw XLSX bytes.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    # Write preamble rows (blank)
    for i in range(header_at):
        ws.append([None])  # blank row
    # Write header row
    ws.append(header_row_cells)
    # Write data rows
    for row in data_rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mock_response(status_code: int, content: bytes) -> MagicMock:
    """Return a MagicMock mimicking a requests.Response."""
    resp = MagicMock()
    resp.status_code = status_code
    resp.content = content
    resp.raise_for_status = MagicMock()
    return resp


# ---------------------------------------------------------------------------
# EMAEparFetcher tests
# ---------------------------------------------------------------------------

class TestEMAEparFetcherValid:
    """Happy-path: well-formed XLSX → success."""

    def test_valid_xlsx_success(self, tmp_path):
        """200 + well-formed XLSX with recognised header → status==success, record_count>=1."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        # Header row has a column name containing an expected token ("Name of medicine")
        header = ["Name of medicine", "Active substance", "Procedure number", "CHMP outcome"]
        data = [["Keytruda", "pembrolizumab", "EMEA/H/C/003820", "Positive"]]
        content = _xlsx_bytes(header, data)

        mock_resp = _mock_response(200, content)

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "success", f"Expected success, got {result}"
        assert result["record_count"] >= 1
        assert isinstance(result["records"], list)
        assert result["hash"] is not None


class TestEMAEparParseError:
    """H4: Corrupt/non-xlsx content → parse_error (NOT source_unavailable)."""

    def test_corrupt_bytes_is_parse_error_not_source_unavailable(self, tmp_path):
        """200 + corrupt bytes → status==parse_error, NOT source_unavailable or failed."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        mock_resp = _mock_response(200, b"definitely not a xlsx zip")

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "parse_error", (
            f"Expected parse_error for corrupt bytes, got {result['status']!r}"
        )
        assert result["status"] != "source_unavailable"
        assert result["status"] != "failed"
        assert result["record_count"] == 0

    def test_ws_none_raises_parse_error(self, tmp_path):
        """Workbook with active=None → status==parse_error (replaces bare assert)."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        # Build valid xlsx bytes (so load_workbook succeeds), then patch wb.active to None
        header = ["Name of medicine", "Active substance"]
        content = _xlsx_bytes(header, [["Keytruda", "pembrolizumab"]])
        mock_resp = _mock_response(200, content)

        fake_wb = MagicMock()
        fake_wb.active = None
        fake_wb.close = MagicMock()

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            with patch("dk_data.ingestion.fetchers.ema_epar.openpyxl.load_workbook", return_value=fake_wb):
                result = fetcher.fetch()

        assert result["status"] == "parse_error", (
            f"Expected parse_error for ws=None, got {result['status']!r}"
        )


class TestEMAEparSchemaError:
    """H4: Valid XLSX structure but wrong/missing headers or 0 data rows → schema_mismatch."""

    def test_header_drift_is_schema_mismatch(self, tmp_path):
        """200 + valid xlsx but row-8 headers contain no expected token → schema_mismatch."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        # Header cells with no token from _EXPECTED_HEADER_TOKENS
        bad_header = ["wibble", "wobble", "zzz", "unrecognized_column"]
        data = [["a", "b", "c", "d"]]
        content = _xlsx_bytes(bad_header, data)
        mock_resp = _mock_response(200, content)

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "schema_mismatch", (
            f"Expected schema_mismatch for header drift, got {result['status']!r}"
        )
        assert result["record_count"] == 0

    def test_zero_data_rows_is_schema_mismatch(self, tmp_path):
        """200 + valid header but ZERO data rows → schema_mismatch (NOT source_unavailable)."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        header = ["Name of medicine", "Active substance", "Procedure number"]
        content = _xlsx_bytes(header, [])  # no data rows
        mock_resp = _mock_response(200, content)

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "schema_mismatch", (
            f"Expected schema_mismatch for 0 data rows, got {result['status']!r}"
        )
        assert result["status"] != "source_unavailable"
        assert result["record_count"] == 0


class TestEMAEparSourceUnavailable:
    """True server-down: all URLs non-200 → source_unavailable (unchanged behavior)."""

    def test_all_urls_404_is_source_unavailable(self, tmp_path):
        """All URLs return 404 → status==source_unavailable."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        mock_resp = _mock_response(404, b"")

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "source_unavailable", (
            f"Expected source_unavailable for all-404, got {result['status']!r}"
        )


class TestEMAEparH5PrimaryFailure:
    """H5: Primary URL failure must be logged at ERROR level, not just WARNING."""

    def test_primary_url_failure_logged_at_error_before_fallback(self, tmp_path, caplog):
        """Primary URL 404 → ERROR log emitted; fetch succeeds via fallback."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        primary_url = _EPAR_URLS[0][0]

        # Build valid xlsx for the fallback response
        header = ["Name of medicine", "Active substance"]
        data = [["Keytruda", "pembrolizumab"]]
        valid_content = _xlsx_bytes(header, data)

        call_count = 0

        def mock_get(url, timeout=None):
            nonlocal call_count
            call_count += 1
            if url == primary_url:
                return _mock_response(404, b"")
            # All other URLs return valid xlsx
            return _mock_response(200, valid_content)

        with caplog.at_level(logging.ERROR, logger="dk_data.ingestion.fetchers.ema_epar"):
            with patch.object(fetcher.session, "get", side_effect=mock_get):
                result = fetcher.fetch()

        # Fetch still ultimately succeeds via fallback
        assert result["status"] == "success", (
            f"Expected success via fallback, got {result['status']!r}"
        )
        # An ERROR-level record must have been emitted by the ema_epar fetcher logger
        # mentioning the PRIMARY failure — filtered by logger name so this test FAILS
        # if the inline logger.error("... PRIMARY ...") in EMAEparFetcher.fetch() is removed.
        fetcher_errors = [
            r for r in caplog.records
            if r.levelno >= logging.ERROR
            and r.name == "dk_data.ingestion.fetchers.ema_epar"
            and "PRIMARY" in r.getMessage()
        ]
        assert len(fetcher_errors) == 1, (
            f"expected exactly one PRIMARY-failure ERROR from the ema_epar fetcher, "
            f"got: {[(r.levelname, r.name, r.getMessage()) for r in caplog.records]}"
        )

    def test_no_error_on_normal_all_success_run(self, tmp_path, caplog):
        """Control: no ERROR emitted when primary URL succeeds normally."""
        fetcher = EMAEparFetcher(data_dir=str(tmp_path))

        header = ["Name of medicine", "Active substance"]
        data = [["Keytruda", "pembrolizumab"]]
        valid_content = _xlsx_bytes(header, data)

        mock_resp = _mock_response(200, valid_content)

        with caplog.at_level(logging.ERROR, logger="dk_data.ingestion.fetchers.ema_epar"):
            with patch.object(fetcher.session, "get", return_value=mock_resp):
                result = fetcher.fetch()

        assert result["status"] == "success"
        # No ERROR records from the ema_epar logger (log_fetch_result in base is ok)
        epar_error_records = [
            r for r in caplog.records
            if r.levelno >= logging.ERROR and r.name == "dk_data.ingestion.fetchers.ema_epar"
        ]
        assert epar_error_records == [], (
            f"Unexpected ERROR logs on clean run: {[(r.levelname, r.message) for r in epar_error_records]}"
        )


# ---------------------------------------------------------------------------
# EMAMolFetcher parallel tests
# ---------------------------------------------------------------------------

class TestEMAMolParseError:
    """H4 (ema_mol): corrupt content → parse_error."""

    def test_emamol_corrupt_is_parse_error(self, tmp_path):
        """200 + corrupt bytes → status==parse_error for EMAMolFetcher."""
        fetcher = EMAMolFetcher(data_dir=str(tmp_path))

        mock_resp = _mock_response(200, b"not an xlsx file at all")

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "parse_error", (
            f"Expected parse_error for corrupt bytes, got {result['status']!r}"
        )
        assert result["record_count"] == 0


class TestEMAMolSchemaError:
    """H4 (ema_mol): bad header / 0 rows → schema_mismatch."""

    def test_emamol_header_drift_is_schema_mismatch(self, tmp_path):
        """200 + xlsx with unrecognised row-8 header → schema_mismatch."""
        fetcher = EMAMolFetcher(data_dir=str(tmp_path))

        bad_header = ["zzz_unknown", "foo_unknown", "bar_unknown"]
        data = [["x", "y", "z"]]
        content = _xlsx_bytes(bad_header, data)
        mock_resp = _mock_response(200, content)

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "schema_mismatch", (
            f"Expected schema_mismatch for header drift, got {result['status']!r}"
        )

    def test_emamol_zero_rows_is_schema_mismatch(self, tmp_path):
        """200 + valid header but 0 data rows → schema_mismatch."""
        fetcher = EMAMolFetcher(data_dir=str(tmp_path))

        header = ["Category", "Name of medicine", "Active substance"]
        content = _xlsx_bytes(header, [])
        mock_resp = _mock_response(200, content)

        with patch.object(fetcher.session, "get", return_value=mock_resp):
            result = fetcher.fetch()

        assert result["status"] == "schema_mismatch", (
            f"Expected schema_mismatch for 0 data rows, got {result['status']!r}"
        )


class TestEMAMolH5PrimaryFailure:
    """H5 (ema_mol): single URL is the primary; its failure must log at ERROR."""

    def test_emamol_primary_failure_logged_at_error(self, tmp_path, caplog):
        """Primary (only) URL fails → ERROR logged, status==source_unavailable."""
        fetcher = EMAMolFetcher(data_dir=str(tmp_path))

        # Return non-200 to trigger primary failure
        mock_resp = _mock_response(503, b"")

        with caplog.at_level(logging.ERROR, logger="dk_data.ingestion.fetchers.ema_mol"):
            with patch.object(fetcher.session, "get", return_value=mock_resp):
                result = fetcher.fetch()

        assert result["status"] == "source_unavailable"
        # Filter by ema_mol logger name AND "PRIMARY" message so this test FAILS
        # if the inline logger.error("... PRIMARY ...") in EMAMolFetcher.fetch() is removed.
        fetcher_errors = [
            r for r in caplog.records
            if r.levelno >= logging.ERROR
            and r.name == "dk_data.ingestion.fetchers.ema_mol"
            and "PRIMARY" in r.getMessage()
        ]
        assert len(fetcher_errors) == 1, (
            f"expected exactly one PRIMARY-failure ERROR from the ema_mol fetcher, "
            f"got: {[(r.levelname, r.name, r.getMessage()) for r in caplog.records]}"
        )
