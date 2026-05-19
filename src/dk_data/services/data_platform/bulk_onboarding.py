"""
Bulk Onboarding Service

Handles bulk data import from CSV/Excel files with async processing
and progress tracking.

Part of DK Molecule Data Platform (012-dk-data-platform)
"""

import asyncio
import csv
import io
import logging
from datetime import datetime
from typing import Dict, Any, Optional, List, AsyncIterator, Tuple
from uuid import UUID, uuid4
from dataclasses import dataclass, field
from enum import Enum
import json

logger = logging.getLogger(__name__)

# Optional Excel support
try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False
    logger.warning("openpyxl not installed, Excel support disabled")


class JobStatus(str, Enum):
    """Bulk upload job status."""
    PENDING = "pending"
    VALIDATING = "validating"
    PROCESSING = "processing"
    COMPLETED = "completed"
    FAILED = "failed"
    CANCELLED = "cancelled"


class RecordStatus(str, Enum):
    """Individual record status."""
    PENDING = "pending"
    VALID = "valid"
    INVALID = "invalid"
    PROCESSED = "processed"
    FAILED = "failed"
    DUPLICATE = "duplicate"


@dataclass
class ValidationError:
    """Validation error for a record."""
    row_number: int
    field: str
    error: str
    value: Any = None


@dataclass
class ProcessingResult:
    """Result of processing a single record."""
    row_number: int
    status: RecordStatus
    molecule_id: Optional[UUID] = None
    errors: List[ValidationError] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


@dataclass
class BulkUploadJob:
    """Bulk upload job tracking."""
    id: UUID
    user_id: UUID
    filename: str
    file_type: str  # csv, xlsx
    status: JobStatus
    total_records: int
    processed_records: int
    successful_records: int
    failed_records: int
    created_at: datetime
    started_at: Optional[datetime] = None
    completed_at: Optional[datetime] = None
    error_message: Optional[str] = None
    metadata: Optional[Dict[str, Any]] = None


@dataclass
class FieldMapping:
    """Mapping configuration for a field."""
    source_column: str
    target_field: str
    required: bool = False
    transform: Optional[str] = None  # Function name for transformation
    default: Optional[Any] = None


# Standard field mappings for molecule data
DEFAULT_MOLECULE_MAPPINGS = {
    "name": FieldMapping("name", "name", required=True),
    "inchi_key": FieldMapping("inchi_key", "inchi_key"),
    "smiles": FieldMapping("smiles", "smiles"),
    "molecular_formula": FieldMapping("molecular_formula", "molecular_formula"),
    "molecular_weight": FieldMapping("molecular_weight", "molecular_weight", transform="to_float"),
    "cas_number": FieldMapping("cas_number", "cas_number"),
    "drugbank_id": FieldMapping("drugbank_id", "external_ids.drugbank"),
    "chembl_id": FieldMapping("chembl_id", "external_ids.chembl"),
    "pubchem_cid": FieldMapping("pubchem_cid", "external_ids.pubchem"),
}


class BulkOnboardingService:
    """
    Service for bulk data onboarding via CSV/Excel uploads.

    Features:
    - CSV and Excel file parsing
    - Configurable field mappings
    - Async processing with progress tracking
    - Validation with detailed error reporting
    - Duplicate detection
    - Batch database operations
    """

    BATCH_SIZE = 100  # Records per batch for DB operations
    MAX_ERRORS_PER_JOB = 1000  # Max errors to track before stopping

    def __init__(self, db_pool, bronze_service=None, entity_resolver=None):
        """
        Initialize bulk onboarding service.

        Args:
            db_pool: Database connection pool
            bronze_service: Optional BronzeIngestionService for record processing
            entity_resolver: Optional EntityResolver for duplicate detection
        """
        self.db_pool = db_pool
        self.bronze_service = bronze_service
        self.entity_resolver = entity_resolver
        self._active_jobs: Dict[UUID, BulkUploadJob] = {}

    async def create_job(
        self,
        user_id: UUID,
        filename: str,
        file_content: bytes,
        mappings: Optional[Dict[str, FieldMapping]] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> BulkUploadJob:
        """
        Create a new bulk upload job.

        Args:
            user_id: ID of the user initiating the upload
            filename: Original filename
            file_content: Raw file content
            mappings: Optional custom field mappings
            metadata: Optional job metadata

        Returns:
            Created BulkUploadJob
        """
        # Detect file type
        file_type = self._detect_file_type(filename, file_content)

        # Count records
        total_records = await self._count_records(file_content, file_type)

        job = BulkUploadJob(
            id=uuid4(),
            user_id=user_id,
            filename=filename,
            file_type=file_type,
            status=JobStatus.PENDING,
            total_records=total_records,
            processed_records=0,
            successful_records=0,
            failed_records=0,
            created_at=datetime.utcnow(),
            metadata=metadata or {},
        )

        # Store in database
        async with self.db_pool.acquire() as conn:
            await conn.execute("""
                INSERT INTO platform.bulk_upload_jobs (
                    id, user_id, filename, file_type, status,
                    total_records, processed_records, successful_records,
                    failed_records, created_at, metadata
                ) VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11)
            """,
                job.id, job.user_id, job.filename, job.file_type,
                job.status.value, job.total_records, job.processed_records,
                job.successful_records, job.failed_records, job.created_at,
                json.dumps(job.metadata) if job.metadata else None
            )

            # Store file content for async processing
            await conn.execute("""
                INSERT INTO platform.bulk_upload_files (job_id, content, mappings)
                VALUES ($1, $2, $3)
            """,
                job.id, file_content,
                json.dumps({k: v.__dict__ for k, v in (mappings or DEFAULT_MOLECULE_MAPPINGS).items()})
            )

        self._active_jobs[job.id] = job
        logger.info(f"Created bulk upload job {job.id} with {total_records} records")

        return job

    async def start_processing(self, job_id: UUID) -> BulkUploadJob:
        """
        Start async processing of a bulk upload job.

        Args:
            job_id: ID of the job to process

        Returns:
            Updated BulkUploadJob
        """
        job = await self.get_job(job_id)
        if not job:
            raise ValueError(f"Job {job_id} not found")

        if job.status != JobStatus.PENDING:
            raise ValueError(f"Job {job_id} is not pending (status: {job.status})")

        # Update status
        job.status = JobStatus.PROCESSING
        job.started_at = datetime.utcnow()
        await self._update_job(job)

        # Start async processing
        asyncio.create_task(self._process_job(job_id))

        return job

    async def _process_job(self, job_id: UUID):
        """Process a bulk upload job asynchronously."""
        try:
            # Load job and file content
            async with self.db_pool.acquire() as conn:
                row = await conn.fetchrow("""
                    SELECT j.*, f.content, f.mappings
                    FROM platform.bulk_upload_jobs j
                    JOIN platform.bulk_upload_files f ON f.job_id = j.id
                    WHERE j.id = $1
                """, job_id)

            if not row:
                logger.error(f"Job {job_id} not found in database")
                return

            file_content = row['content']
            mappings_dict = json.loads(row['mappings']) if row['mappings'] else {}
            mappings = {
                k: FieldMapping(**v) for k, v in mappings_dict.items()
            }

            job = await self.get_job(job_id)

            # Update status to validating
            job.status = JobStatus.VALIDATING
            await self._update_job(job)

            # Parse and validate records
            records = []
            errors = []

            async for row_num, record, row_errors in self._parse_file(
                file_content, row['file_type'], mappings
            ):
                if row_errors:
                    errors.extend(row_errors)
                    job.failed_records += 1
                else:
                    records.append((row_num, record))

                # Store errors in database
                if row_errors:
                    await self._store_errors(job_id, row_errors)

                # Check error limit
                if len(errors) >= self.MAX_ERRORS_PER_JOB:
                    job.status = JobStatus.FAILED
                    job.error_message = f"Too many validation errors ({len(errors)})"
                    await self._update_job(job)
                    return

            # Update to processing
            job.status = JobStatus.PROCESSING
            await self._update_job(job)

            # Process in batches
            for i in range(0, len(records), self.BATCH_SIZE):
                batch = records[i:i + self.BATCH_SIZE]
                results = await self._process_batch(job_id, batch)

                for result in results:
                    job.processed_records += 1
                    if result.status == RecordStatus.PROCESSED:
                        job.successful_records += 1
                    else:
                        job.failed_records += 1

                await self._update_job(job)

            # Complete job
            job.status = JobStatus.COMPLETED
            job.completed_at = datetime.utcnow()
            await self._update_job(job)

            logger.info(
                f"Completed job {job_id}: {job.successful_records} successful, "
                f"{job.failed_records} failed"
            )

        except Exception as e:
            logger.exception(f"Error processing job {job_id}")
            job = await self.get_job(job_id)
            if job:
                job.status = JobStatus.FAILED
                job.error_message = str(e)
                job.completed_at = datetime.utcnow()
                await self._update_job(job)

    async def _process_batch(
        self,
        job_id: UUID,
        batch: List[Tuple[int, Dict[str, Any]]]
    ) -> List[ProcessingResult]:
        """Process a batch of records."""
        results = []

        async with self.db_pool.acquire() as conn:
            for row_num, record in batch:
                try:
                    # Check for duplicates if entity resolver available
                    if self.entity_resolver and record.get("inchi_key"):
                        existing = await self.entity_resolver.find_by_inchi_key(
                            record["inchi_key"]
                        )
                        if existing:
                            results.append(ProcessingResult(
                                row_number=row_num,
                                status=RecordStatus.DUPLICATE,
                                molecule_id=existing.id,
                                warnings=[f"Duplicate of existing molecule {existing.id}"]
                            ))
                            continue

                    # Insert into bronze layer
                    if self.bronze_service:
                        molecule_id = await self.bronze_service.ingest_record(
                            source="bulk_upload",
                            record=record,
                            metadata={"job_id": str(job_id), "row": row_num}
                        )
                    else:
                        # Direct insert if no bronze service
                        molecule_id = uuid4()
                        await conn.execute("""
                            INSERT INTO mol_bronze.molecules (
                                id, source, raw_data, ingested_at
                            ) VALUES ($1, $2, $3, NOW())
                        """, molecule_id, "bulk_upload", json.dumps(record))

                    results.append(ProcessingResult(
                        row_number=row_num,
                        status=RecordStatus.PROCESSED,
                        molecule_id=molecule_id
                    ))

                except Exception as e:
                    logger.error(f"Error processing row {row_num}: {e}")
                    results.append(ProcessingResult(
                        row_number=row_num,
                        status=RecordStatus.FAILED,
                        errors=[ValidationError(
                            row_number=row_num,
                            field="*",
                            error=str(e)
                        )]
                    ))

        return results

    async def _parse_file(
        self,
        content: bytes,
        file_type: str,
        mappings: Dict[str, FieldMapping]
    ) -> AsyncIterator[Tuple[int, Dict[str, Any], List[ValidationError]]]:
        """Parse file content and yield records with validation."""
        if file_type == "csv":
            async for item in self._parse_csv(content, mappings):
                yield item
        elif file_type == "xlsx":
            async for item in self._parse_excel(content, mappings):
                yield item
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    async def _parse_csv(
        self,
        content: bytes,
        mappings: Dict[str, FieldMapping]
    ) -> AsyncIterator[Tuple[int, Dict[str, Any], List[ValidationError]]]:
        """Parse CSV content."""
        # Decode content
        text = content.decode('utf-8-sig')  # Handle BOM
        reader = csv.DictReader(io.StringIO(text))

        for row_num, row in enumerate(reader, start=2):  # Start at 2 (header is 1)
            record, errors = self._map_and_validate_row(row_num, row, mappings)
            yield row_num, record, errors
            await asyncio.sleep(0)  # Yield control

    async def _parse_excel(
        self,
        content: bytes,
        mappings: Dict[str, FieldMapping]
    ) -> AsyncIterator[Tuple[int, Dict[str, Any], List[ValidationError]]]:
        """Parse Excel content."""
        if not EXCEL_SUPPORT:
            raise ValueError("Excel support requires openpyxl library")

        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        sheet = workbook.active

        # Get headers from first row
        headers = [cell.value for cell in sheet[1]]

        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row_dict = dict(zip(headers, row))
            record, errors = self._map_and_validate_row(row_num, row_dict, mappings)
            yield row_num, record, errors
            await asyncio.sleep(0)  # Yield control

        workbook.close()

    def _map_and_validate_row(
        self,
        row_num: int,
        row: Dict[str, Any],
        mappings: Dict[str, FieldMapping]
    ) -> Tuple[Dict[str, Any], List[ValidationError]]:
        """Map row data to target fields and validate."""
        record = {}
        errors = []

        for field_name, mapping in mappings.items():
            value = row.get(mapping.source_column)

            # Handle required fields
            if mapping.required and (value is None or value == ""):
                errors.append(ValidationError(
                    row_number=row_num,
                    field=mapping.source_column,
                    error=f"Required field '{mapping.source_column}' is missing",
                    value=value
                ))
                continue

            # Apply default if value is missing
            if value is None or value == "":
                value = mapping.default

            # Apply transformation
            if value is not None and mapping.transform:
                try:
                    value = self._apply_transform(value, mapping.transform)
                except Exception as e:
                    errors.append(ValidationError(
                        row_number=row_num,
                        field=mapping.source_column,
                        error=f"Transform '{mapping.transform}' failed: {e}",
                        value=value
                    ))
                    continue

            # Handle nested fields (e.g., "external_ids.drugbank")
            if "." in mapping.target_field:
                parts = mapping.target_field.split(".")
                current = record
                for part in parts[:-1]:
                    if part not in current:
                        current[part] = {}
                    current = current[part]
                current[parts[-1]] = value
            else:
                record[mapping.target_field] = value

        return record, errors

    def _apply_transform(self, value: Any, transform: str) -> Any:
        """Apply transformation to a value."""
        transforms = {
            "to_float": lambda v: float(v) if v else None,
            "to_int": lambda v: int(v) if v else None,
            "to_upper": lambda v: str(v).upper() if v else None,
            "to_lower": lambda v: str(v).lower() if v else None,
            "strip": lambda v: str(v).strip() if v else None,
            "to_bool": lambda v: str(v).lower() in ("true", "yes", "1") if v else False,
        }

        if transform not in transforms:
            raise ValueError(f"Unknown transform: {transform}")

        return transforms[transform](value)

    def _detect_file_type(self, filename: str, content: bytes) -> str:
        """Detect file type from filename and content."""
        filename_lower = filename.lower()

        if filename_lower.endswith(".csv"):
            return "csv"
        elif filename_lower.endswith((".xlsx", ".xls")):
            return "xlsx"

        # Check magic bytes
        if content[:4] == b'PK\x03\x04':  # ZIP (xlsx)
            return "xlsx"

        # Default to CSV
        return "csv"

    async def _count_records(self, content: bytes, file_type: str) -> int:
        """Count records in file."""
        if file_type == "csv":
            text = content.decode('utf-8-sig')
            return sum(1 for _ in csv.reader(io.StringIO(text))) - 1  # Subtract header
        elif file_type == "xlsx" and EXCEL_SUPPORT:
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            sheet = workbook.active
            count = sum(1 for _ in sheet.iter_rows(min_row=2))
            workbook.close()
            return count
        return 0

    async def _update_job(self, job: BulkUploadJob):
        """Update job in database."""
        async with self.db_pool.acquire() as conn:
            await conn.execute("""
                UPDATE platform.bulk_upload_jobs
                SET status = $2, processed_records = $3, successful_records = $4,
                    failed_records = $5, started_at = $6, completed_at = $7,
                    error_message = $8
                WHERE id = $1
            """,
                job.id, job.status.value, job.processed_records,
                job.successful_records, job.failed_records, job.started_at,
                job.completed_at, job.error_message
            )

        self._active_jobs[job.id] = job

    async def _store_errors(self, job_id: UUID, errors: List[ValidationError]):
        """Store validation errors in database."""
        async with self.db_pool.acquire() as conn:
            for error in errors:
                await conn.execute("""
                    INSERT INTO platform.bulk_upload_errors (
                        id, job_id, row_number, field, error, value, created_at
                    ) VALUES ($1, $2, $3, $4, $5, $6, NOW())
                """,
                    uuid4(), job_id, error.row_number, error.field,
                    error.error, str(error.value) if error.value else None
                )

    async def get_job(self, job_id: UUID) -> Optional[BulkUploadJob]:
        """Get job by ID."""
        # Check cache first
        if job_id in self._active_jobs:
            return self._active_jobs[job_id]

        # Load from database
        async with self.db_pool.acquire() as conn:
            row = await conn.fetchrow("""
                SELECT * FROM platform.bulk_upload_jobs WHERE id = $1
            """, job_id)

        if not row:
            return None

        return BulkUploadJob(
            id=row['id'],
            user_id=row['user_id'],
            filename=row['filename'],
            file_type=row['file_type'],
            status=JobStatus(row['status']),
            total_records=row['total_records'],
            processed_records=row['processed_records'],
            successful_records=row['successful_records'],
            failed_records=row['failed_records'],
            created_at=row['created_at'],
            started_at=row['started_at'],
            completed_at=row['completed_at'],
            error_message=row['error_message'],
            metadata=json.loads(row['metadata']) if row['metadata'] else None,
        )

    async def get_job_errors(
        self,
        job_id: UUID,
        limit: int = 100,
        offset: int = 0
    ) -> List[Dict[str, Any]]:
        """Get errors for a job."""
        async with self.db_pool.acquire() as conn:
            rows = await conn.fetch("""
                SELECT row_number, field, error, value, created_at
                FROM platform.bulk_upload_errors
                WHERE job_id = $1
                ORDER BY row_number
                LIMIT $2 OFFSET $3
            """, job_id, limit, offset)

        return [
            {
                "row_number": r['row_number'],
                "field": r['field'],
                "error": r['error'],
                "value": r['value'],
                "created_at": r['created_at'].isoformat() if r['created_at'] else None,
            }
            for r in rows
        ]

    async def get_job_progress(self, job_id: UUID) -> Dict[str, Any]:
        """Get job progress summary."""
        job = await self.get_job(job_id)
        if not job:
            return {"error": "Job not found"}

        progress = 0
        if job.total_records > 0:
            progress = (job.processed_records / job.total_records) * 100

        return {
            "job_id": str(job.id),
            "status": job.status.value,
            "progress_percent": round(progress, 1),
            "total_records": job.total_records,
            "processed_records": job.processed_records,
            "successful_records": job.successful_records,
            "failed_records": job.failed_records,
            "started_at": job.started_at.isoformat() if job.started_at else None,
            "completed_at": job.completed_at.isoformat() if job.completed_at else None,
            "error_message": job.error_message,
        }

    async def cancel_job(self, job_id: UUID) -> bool:
        """Cancel a running job."""
        job = await self.get_job(job_id)
        if not job:
            return False

        if job.status not in (JobStatus.PENDING, JobStatus.PROCESSING, JobStatus.VALIDATING):
            return False

        job.status = JobStatus.CANCELLED
        job.completed_at = datetime.utcnow()
        await self._update_job(job)

        logger.info(f"Cancelled job {job_id}")
        return True

    async def list_jobs(
        self,
        user_id: Optional[UUID] = None,
        status: Optional[JobStatus] = None,
        limit: int = 20,
        offset: int = 0
    ) -> List[Dict[str, Any]]:
        """List bulk upload jobs."""
        async with self.db_pool.acquire() as conn:
            query = "SELECT * FROM platform.bulk_upload_jobs WHERE 1=1"
            params = []
            param_idx = 1

            if user_id:
                query += f" AND user_id = ${param_idx}"
                params.append(user_id)
                param_idx += 1

            if status:
                query += f" AND status = ${param_idx}"
                params.append(status.value)
                param_idx += 1

            query += f" ORDER BY created_at DESC LIMIT ${param_idx} OFFSET ${param_idx + 1}"
            params.extend([limit, offset])

            rows = await conn.fetch(query, *params)

        return [
            {
                "id": str(r['id']),
                "user_id": str(r['user_id']),
                "filename": r['filename'],
                "file_type": r['file_type'],
                "status": r['status'],
                "total_records": r['total_records'],
                "processed_records": r['processed_records'],
                "successful_records": r['successful_records'],
                "failed_records": r['failed_records'],
                "created_at": r['created_at'].isoformat() if r['created_at'] else None,
                "completed_at": r['completed_at'].isoformat() if r['completed_at'] else None,
            }
            for r in rows
        ]

    async def get_template(self, format: str = "csv") -> bytes:
        """Get a template file for bulk upload."""
        headers = [
            "name", "inchi_key", "smiles", "molecular_formula",
            "molecular_weight", "cas_number", "drugbank_id",
            "chembl_id", "pubchem_cid"
        ]

        if format == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(headers)
            # Add example row
            writer.writerow([
                "Aspirin", "BSYNRYMUTXBXSQ-UHFFFAOYSA-N",
                "CC(=O)OC1=CC=CC=C1C(O)=O", "C9H8O4",
                "180.158", "50-78-2", "DB00945",
                "CHEMBL25", "2244"
            ])
            return output.getvalue().encode('utf-8')

        elif format == "xlsx" and EXCEL_SUPPORT:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Molecules"

            # Headers
            for col, header in enumerate(headers, start=1):
                sheet.cell(row=1, column=col, value=header)

            # Example row
            example = [
                "Aspirin", "BSYNRYMUTXBXSQ-UHFFFAOYSA-N",
                "CC(=O)OC1=CC=CC=C1C(O)=O", "C9H8O4",
                180.158, "50-78-2", "DB00945",
                "CHEMBL25", "2244"
            ]
            for col, value in enumerate(example, start=1):
                sheet.cell(row=2, column=col, value=value)

            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue()

        else:
            raise ValueError(f"Unsupported format: {format}")
